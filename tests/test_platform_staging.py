"""Staging a PLATFORM recording so the LOCAL worker can execute it.

The worker only ever downloads from this machine's bucket, so a recording that
lives on the platform is copied here first. These tests pin the two properties
that make the copy trustworthy: the workbook is stored byte for byte, and the
run path still finds the repeatable rows in a workbook this UI did not write.
"""

from __future__ import annotations

import base64
import importlib.util
import io
import sys
from pathlib import Path


_APP_PATH = Path(__file__).resolve().parents[1] / "app.py"
_SPEC = importlib.util.spec_from_file_location("agent_shubham_app", _APP_PATH)
assert _SPEC and _SPEC.loader
app = importlib.util.module_from_spec(_SPEC)
# Registered before exec: app.py's pydantic models carry string annotations, and
# pydantic resolves them through sys.modules[cls.__module__]. Without this,
# building an UploadBody raises "is not fully defined".
sys.modules[_SPEC.name] = app
_SPEC.loader.exec_module(app)


class _FakeS3:
    """Captures put_object calls; get_object serves whatever was put."""

    def __init__(self) -> None:
        self.objects: dict[str, bytes] = {}

    def put_object(self, *, Bucket: str, Key: str, Body: bytes, ContentType: str = "") -> None:
        self.objects[Key] = Body

    def get_object(self, *, Bucket: str, Key: str):
        if Key not in self.objects:
            raise KeyError(Key)
        return {"Body": io.BytesIO(self.objects[Key])}


def _workbook(sheets: dict[str, list[list[str]]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        sheet = wb.create_sheet(title=title)
        for row in rows:
            sheet.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- verbatim workbook
def test_upload_stores_the_workbook_verbatim(monkeypatch) -> None:
    """params_b64 is written unchanged -- not rebuilt from the params rows.

    The rebuild path is not lossless (it injects a ref_id column, drops blank
    cells from the repeatable sheet, and drops sheets the parser does not read),
    so a staged recording is COPIED rather than regenerated.
    """
    fake = _FakeS3()
    monkeypatch.setattr(app, "_s3", lambda: fake)
    upstream = _workbook({"params": [["invoice_number", "memo"], ["INV-1", ""]]})

    result = app.upload(
        app.UploadBody(
            name="demo",
            script="page.goto('https://example.test')",
            params={"params": []},
            params_b64=base64.b64encode(upstream).decode(),
            bucket="tenant",
            register=False,
        )
    )

    assert result["params_verbatim"] is True
    assert fake.objects["recordings/demo/demo_params.xlsx"] == upstream
    # Byte-identical means the header row is untouched: same columns, same
    # order, and no injected ref_id.
    import openpyxl

    stored = openpyxl.load_workbook(io.BytesIO(fake.objects["recordings/demo/demo_params.xlsx"]))
    assert [c.value for c in stored["params"][1]] == ["invoice_number", "memo"]


def test_upload_without_params_b64_still_builds_the_workbook(monkeypatch) -> None:
    fake = _FakeS3()
    monkeypatch.setattr(app, "_s3", lambda: fake)

    result = app.upload(
        app.UploadBody(
            name="demo",
            script="page.goto('https://example.test')",
            params={"params": [{"invoice_number": "INV-1"}]},
            bucket="tenant",
            register=False,
        )
    )

    assert result["params_verbatim"] is False
    # Rebuilt from the rows -- and the writer adds its own ref_id column, which
    # is exactly the kind of difference staging avoids by copying.
    assert app._parse_params_back(
        fake.objects["recordings/demo/demo_params.xlsx"], "_params.xlsx"
    ) == [{"ref_id": "1", "invoice_number": "INV-1"}]


def test_upload_rejects_unusable_params_b64(monkeypatch) -> None:
    monkeypatch.setattr(app, "_s3", lambda: _FakeS3())

    try:
        app.upload(
            app.UploadBody(
                name="demo",
                script="page.goto('https://example.test')",
                params={"params": []},
                params_b64="not base64!!",
                bucket="tenant",
                register=False,
            )
        )
    except app.HTTPException as exc:
        assert exc.status_code == 400
        assert "params_b64" in str(exc.detail)
    else:  # pragma: no cover - defensive
        raise AssertionError("expected HTTPException")


def test_register_false_skips_the_recorded_flows_row(monkeypatch) -> None:
    fake = _FakeS3()
    monkeypatch.setattr(app, "_s3", lambda: fake)
    calls: list[str] = []
    monkeypatch.setattr(
        app, "upsert_recorded_flow", lambda **kwargs: calls.append(kwargs["name"]) or {}
    )
    body = dict(
        name="demo",
        script="page.goto('https://example.test')",
        params={"params": [{"a": "1"}]},
        bucket="tenant",
    )

    staged = app.upload(app.UploadBody(**body, register=False))
    assert staged["registered"] is False
    assert calls == []

    saved = app.upload(app.UploadBody(**body, register=True))
    assert saved["registered"] is True
    assert calls == ["demo"]


# --------------------------------------------------------------------------- foreign sheet names
def test_repeatable_sheet_in_prefers_this_uis_sheet() -> None:
    raw = _workbook(
        {
            "params": [["a"], ["1"]],
            "context": [["k"], ["v"]],
            "multi_line": [["qty"], ["2"]],
            "line_items": [["qty"], ["3"]],
        }
    )
    assert app._repeatable_sheet_in(raw) == "line_items"


def test_repeatable_sheet_in_finds_a_foreign_sheet() -> None:
    # What the ingest agent writes: params + context + multi_line.
    raw = _workbook(
        {"params": [["a"], ["1"]], "context": [["k"], ["v"]], "multi_line": [["qty"], ["2"]]}
    )
    assert app._repeatable_sheet_in(raw) == "multi_line"


def test_repeatable_sheet_in_reports_none_when_there_is_none() -> None:
    raw = _workbook({"params": [["a"], ["1"]], "context": [["k"], ["v"]]})
    assert app._repeatable_sheet_in(raw) == ""


def test_parse_params_reads_a_foreign_repeatable_sheet() -> None:
    raw = _workbook(
        {
            "params": [["invoice_number"], ["INV-1"]],
            "multi_line": [["ref_id", "quantity"], ["INV-1", "2"]],
        }
    )

    out = app.parse_params(
        app.ParseParamsBody(filename="demo_params.xlsx", content_b64=base64.b64encode(raw).decode())
    )

    assert out["params"] == [{"invoice_number": "INV-1"}]
    assert out["multi_line"] == [{"ref_id": "INV-1", "quantity": "2"}]
    # Reported, so an empty list can never be mistaken for "no line items".
    assert out["multi_line_sheet"] == "multi_line"


def test_saved_payload_finds_repeatable_rows_without_a_sidecar(monkeypatch) -> None:
    """A staged recording has no recording_config naming its repeatable sheet.

    Falling back to the sheet that is actually there is what stops the run from
    executing the repeated block with zero rows and still reporting green.
    """
    fake = _FakeS3()
    fake.objects["recordings/demo/demo_params.xlsx"] = _workbook(
        {
            "params": [["invoice_number"], ["INV-1"]],
            "multi_line": [["ref_id", "quantity"], ["INV-1", "2"]],
        }
    )
    monkeypatch.setattr(app, "_s3", lambda: fake)
    monkeypatch.setattr(app, "_load_recording_config", lambda name, bucket: {})

    params_rows, multi_line_rows, params_key = app._load_saved_runtime_payload("demo", "tenant")

    assert params_rows == [{"invoice_number": "INV-1"}]
    assert multi_line_rows == [{"ref_id": "INV-1", "quantity": "2"}]
    assert params_key == "recordings/demo/demo_params.xlsx"
